# -*- coding: utf-8 -*-
"""问题二、三的守恒核查与问题三边界延拓敏感性分析。"""
from __future__ import annotations

import json
from pathlib import Path
import sys

import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src" / "q2_full"))
sys.path.insert(0, str(ROOT / "src" / "q3_drying"))
import solver_q2 as q2
import solver_q3 as q3


def cumulative_mass_balance(path: Path, extension_end: float | None = None):
    """用题目要求的0.1 cm输出剖面独立复核累计水分收支。"""
    data = np.load(path)
    t = data["t"]
    c = data["C"]
    radius = np.linspace(0.0, q2.RADIUS, c.shape[1])
    inventory = np.trapezoid(c * radius, radius, axis=1)
    ambient = q2.env_moisture(t)
    ambient[0] = 0.01963
    if extension_end is not None:
        ambient[t > extension_end] = float(q2.env_moisture(extension_end))
    outward_flux = q2.H_M * q2.RADIUS * (c[:, -1] - ambient)
    discharged = np.zeros_like(t)
    discharged[1:] = np.cumsum(
        0.5 * (outward_flux[1:] + outward_flux[:-1]) * np.diff(t)
    )
    residual = inventory - inventory[0] + discharged
    total_loss = inventory[0] - inventory[-1]
    return {
        "initial_inventory": float(inventory[0]),
        "total_loss": float(total_loss),
        "final_relative_residual": float(abs(residual[-1]) / total_loss),
        "maximum_relative_residual": float(np.max(np.abs(residual)) / total_loss),
    }


def plateau_noise_stats(path: Path):
    ws = load_workbook(path, data_only=True, read_only=True).active
    rows = [row for row in ws.iter_rows(min_row=2, values_only=True) if row[0] is not None]
    data = np.asarray(rows, dtype=float)
    plateau = data[(data[:, 0] >= q2.T_SWITCH) & (data[:, 0] <= q3.BOUNDARY_DATA_END)]
    return {
        "temperature_mean_c": float(np.mean(plateau[:, 1])),
        "temperature_sd_c": float(np.std(plateau[:, 1], ddof=1)),
        "moisture_mean": float(np.mean(plateau[:, 2])),
        "moisture_sd": float(np.std(plateau[:, 2], ddof=1)),
    }


def sensitivity(stats):
    t_delta = 3.0 * stats["temperature_sd_c"]
    c_delta = 3.0 * stats["moisture_sd"]
    cases = [
        ("有利延拓", t_delta, -c_delta),
        ("基准延拓", 0.0, 0.0),
        ("不利延拓", -t_delta, c_delta),
    ]
    results = []
    for name, dt_c, dc in cases:
        _, _, _, crossing, _ = q3.run(
            refine=1,
            dt=10.0,
            temperature_extension_offset=dt_c,
            moisture_extension_offset=dc,
        )
        results.append({
            "case": name,
            "temperature_offset_c": float(dt_c),
            "moisture_offset": float(dc),
            "drying_time_s": float(crossing),
            "drying_time_h": float(crossing / 3600.0),
        })
    baseline = next(row["drying_time_h"] for row in results if row["case"] == "基准延拓")
    for row in results:
        row["relative_change"] = float((row["drying_time_h"] - baseline) / baseline)
    return results


def main():
    q2_path = ROOT / "outputs/q2/q2_fullA3_dt05_k1.npz"
    q3_path = ROOT / "outputs/q3/q3_4hfixed_dt5_k1.npz"
    if not q2_path.exists():
        q2_path.parent.mkdir(parents=True, exist_ok=True)
        t, temperature, moisture, max_iterations = q2.run(refine=1, dt=0.5)
        np.savez(q2_path, t=t, T=temperature, C=moisture,
                 max_iterations=max_iterations)
    if not q3_path.exists():
        q3_path.parent.mkdir(parents=True, exist_ok=True)
        t, temperature, moisture, crossing, max_iterations = q3.run(refine=1, dt=5.0)
        np.savez(q3_path, t=t, T=temperature, C=moisture,
                 crossing=crossing, max_iterations=max_iterations)

    stats = plateau_noise_stats(ROOT / "data" / "raw" / "附件1.xlsx")
    result = {
        "mass_balance": {
            "problem2": cumulative_mass_balance(q2_path),
            "problem3": cumulative_mass_balance(
                q3_path, q3.BOUNDARY_DATA_END
            ),
        },
        "plateau_noise": stats,
        "sensitivity": sensitivity(stats),
    }
    out = ROOT / "outputs/q3/q3_audit.json"
    out.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
