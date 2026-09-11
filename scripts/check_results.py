import numpy as np
import openpyxl

print('='*70)
print('问题1 结果精度检查')
print('='*70)

# 1. 检查最终结果文件
print('\n[1. 输出文件检查]')
result_files = ['outputs/q1/result1_final.xlsx', 'outputs/q1/result1_v2.xlsx']
for f in result_files:
    try:
        wb = openpyxl.load_workbook(f, data_only=True)
        print(f'\nOK {f}')
        print(f'  工作表: {wb.sheetnames}')
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            print(f'  - {sheet_name}: {ws.max_row}行 x {ws.max_column}列')
    except Exception as e:
        print(f'FAIL {f}: {e}')

# 2. 读取最新的npz文件，检查数值精度
print('\n[2. 数值精度检查]')
data = np.load('outputs/q1/q1_v2.npz')
t = data['t']
r = data['r']
T = data['T']
C = data['C']

print(f'\n时间步: {len(t)} 个 (0 - {t[-1]:.0f}秒)')
print(f'空间网格: {len(r)} 个节点 (dr = {r[1]-r[0]:.4f} m)')
print(f'输出形状: T{T.shape}, C{C.shape}')

# 检查题目要求的时间点和位置
required_times = [100, 300, 600, 900, 1200, 1500, 1800]
required_r = [0, 0.005, 0.01, 0.015, 0.02]  # 0, 0.5, 1, 1.5, 2 cm

print(f'\n[3. 题目要求的输出点检查]')
print(f'\n温度分布(单位: C)')
print('time(s) | r=0.0cm | r=0.5cm | r=1.0cm | r=1.5cm | r=2.0cm')
print('-'*70)

for t_req in required_times:
    idx_t = int(t_req)
    r_indices = [np.argmin(np.abs(r - ri)) for ri in required_r]
    temps = [T[idx_t, ri] for ri in r_indices]
    print(f'{t_req:4d}    | ' + ' | '.join([f'{temp:7.3f}' for temp in temps]))

print(f'\n水分浓度分布(单位: kg/kg)')
print('time(s) | r=0.0cm | r=0.5cm | r=1.0cm | r=1.5cm | r=2.0cm')
print('-'*70)

for t_req in required_times:
    idx_t = int(t_req)
    r_indices = [np.argmin(np.abs(r - ri)) for ri in required_r]
    moistures = [C[idx_t, ri] for ri in r_indices]
    print(f'{t_req:4d}    | ' + ' | '.join([f'{m:7.4f}' for m in moistures]))

# 4. 物理合理性检查
print(f'\n[4. 物理合理性检查]')

T_final = T[-1, :]
print(f'\n温度分布(1800s):')
print(f'  表面 T(R) = {T_final[-1]:.3f}C')
print(f'  中心 T(0) = {T_final[0]:.3f}C')
gradient_ok = T_final[-1] > T_final[0]
print(f'  梯度: {"OK 正确" if gradient_ok else "FAIL 错误"}(表面应最热)')

C_final = C[-1, :]
print(f'\n水分分布(1800s):')
print(f'  表面 C(R) = {C_final[-1]:.5f} kg/kg')
print(f'  中心 C(0) = {C_final[0]:.5f} kg/kg')
moisture_ok = C_final[0] > C_final[-1]
print(f'  梯度: {"OK 正确" if moisture_ok else "FAIL 错误"}(中心水分应更高)')

print(f'\n初值检查(t=0):')
print(f'  温度 T(r,0) = {T[0, 0]:.2f}C (题面: 28C)')
print(f'  水分 C(r,0) = {C[0, 0]:.3f} kg/kg (题面: 2.55 kg/kg)')
init_T_ok = abs(T[0, 0] - 28.0) < 0.1
init_C_ok = abs(C[0, 0] - 2.55) < 0.01
print(f'  温度初值: {"OK" if init_T_ok else "FAIL"}')
print(f'  水分初值: {"OK" if init_C_ok else "FAIL"}')

total_C_0 = np.sum(C[0, :])
total_C_final = np.sum(C[-1, :])
reduction = (1 - total_C_final/total_C_0)*100
print(f'\n总水分变化:')
print(f'  初始: {total_C_0:.3f}')
print(f'  最终: {total_C_final:.3f}')
print(f'  减少: {reduction:.1f}%')
print(f'  {"OK 合理" if total_C_final < total_C_0 else "FAIL 异常"}(应该减少)')

# 5. 数值稳定性检查
print(f'\n[5. 数值稳定性检查]')

has_nan_T = np.any(np.isnan(T)) or np.any(np.isinf(T))
has_nan_C = np.any(np.isnan(C)) or np.any(np.isinf(C))
has_neg_T = np.any(T < 0)
has_neg_C = np.any(C < 0)

print(f'  温度场: {"FAIL 异常" if has_nan_T else "OK 无NaN/Inf"}')
print(f'  水分场: {"FAIL 异常" if has_nan_C else "OK 无NaN/Inf"}')
print(f'  温度负值: {"FAIL 异常" if has_neg_T else "OK 无"}')
print(f'  水分负值: {"FAIL 异常" if has_neg_C else "OK 无"}')

print(f'\n数值范围:')
print(f'  温度: [{T.min():.2f}, {T.max():.2f}]C')
print(f'  水分: [{C.min():.5f}, {C.max():.5f}] kg/kg')

# 6. 精度评估
print(f'\n[6. 求解精度评估]')

T_change = np.max(np.abs(T[-1, :] - T[-100, :]))
C_change = np.max(np.abs(C[-1, :] - C[-100, :]))

print(f'  温度变化(1700-1800s): {T_change:.6f}C')
print(f'  水分变化(1700-1800s): {C_change:.8f} kg/kg')

steady = T_change < 0.01 and C_change < 1e-5
print(f'  状态: {"OK 已接近稳态" if steady else "WARNING 仍在变化中"}')

# 7. 总体评分
print(f'\n[7. 总体评分]')
scores = {
    '文件完整性': True,
    '初值正确': init_T_ok and init_C_ok,
    '物理梯度': gradient_ok and moisture_ok,
    '数值稳定': not (has_nan_T or has_nan_C or has_neg_T or has_neg_C),
    '水分守恒': total_C_final < total_C_0,
    '接近稳态': steady
}

passed = sum(scores.values())
total = len(scores)

print(f'\n检查项通过: {passed}/{total}')
for key, value in scores.items():
    status = "PASS" if value else "FAIL"
    print(f'  [{status}] {key}')

if passed == total:
    print(f'\n*** 结论: 问题1求解结果 优秀 ***')
elif passed >= total - 1:
    print(f'\n*** 结论: 问题1求解结果 良好 ***')
else:
    print(f'\n*** 结论: 问题1求解结果 需要改进 ***')

print('\n' + '='*70)
print('检查完成')
print('='*70)
