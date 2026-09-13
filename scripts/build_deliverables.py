# -*- coding: utf-8 -*-
"""从题目附件一次性生成 result1.xlsx 至 result4.xlsx。"""
from __future__ import annotations

from pathlib import Path
import sys

import numpy as np
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
for directory in ("q1_preheat", "q2_full", "q3_drying", "q4_shrinkage"):
    sys.path.insert(0, str(ROOT / "src" / directory))

import solver_q1 as q1
import solver_q2 as q2
import solver_q3 as q3
import solver_q4 as q4

DISTANCES_CM = np.arange(0.0, 2.01, 0.1)


def write_two_sheet(path: Path, times, temperature, moisture, include_zero: bool):
    wb = Workbook()
    wb.remove(wb.active)
    start = 0 if include_zero else 1
    for name, field in (("温度", temperature), ("干基含水率", moisture)):
        ws = wb.create_sheet(name)
        ws.append(["时间/s\\距中心距离/cm"] + [round(x, 1) for x in DISTANCES_CM])
        for i in range(start, len(times)):
            ws.append([int(times[i])] + [round(float(v), 4) for v in field[i]])
        for row in ws.iter_rows(min_row=2, min_col=2):
            for cell in row:
                cell.number_format = "0.0000"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_moisture(path: Path, times, moisture, radius=None, xi=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "干基含水率"
    header = ["时间/s\\距中心距离/cm"] + [round(x, 1) for x in DISTANCES_CM]
    if radius is not None:
        header += ["实际半径/cm", "表面干基含水率"]
    ws.append(header)
    for i, time in enumerate(times):
        if radius is None:
            row = [int(time)] + [round(float(v), 4) for v in moisture[i]]
        else:
            profile = q4.interpolate_physical_profile(xi, moisture[i], radius[i], DISTANCES_CM)
            values = [None if np.isnan(v) else round(float(v), 4) for v in profile]
            row = [int(time)] + values + [round(float(radius[i] * 100), 4), round(float(moisture[i, -1]), 4)]
        ws.append(row)
    for row in ws.iter_rows(min_row=2, min_col=2):
        for cell in row:
            cell.number_format = "0.0000"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main():
    out = ROOT / "outputs"

    t1, _, temp1, moisture1, _ = q1.run(k=1, dt=0.25)
    write_two_sheet(out / "q1/result1.xlsx", t1, temp1, moisture1, include_zero=False)

    t2, temp2, moisture2, _ = q2.run(refine=1, dt=0.5)
    write_two_sheet(out / "q2/result2.xlsx", t2, temp2, moisture2, include_zero=False)

    t3, _, moisture3, crossing3, _ = q3.run(refine=1, dt=5.0)
    write_moisture(out / "q3/result3.xlsx", t3, moisture3)

    result4 = q4.run(refine=1, dt=5.0)
    write_moisture(out / "q4/result4.xlsx", result4["t"], result4["C"],
                   radius=result4["R"], xi=result4["xi"])

    print(f"问题三阈值插值时刻：{crossing3:.2f} s ({crossing3 / 3600:.4f} h)")
    print(f"问题四阈值插值时刻：{result4['crossing']:.2f} s ({result4['crossing'] / 3600:.4f} h)")
    print("已生成 result1.xlsx 至 result4.xlsx")


if __name__ == "__main__":
    main()
