# -*- coding: utf-8 -*-
"""问题四：附录4物性、附件2收缩半径下的一维径向温湿耦合求解器。

坐标采用 xi=r/R(t)。求解器在 xi 网格上推进，输出阶段再映射回物理半径 r。
烘房温度在 7200 s 后进入恒温平台；附件1只覆盖 0--14400 s，此后将烘房
水分浓度固定为 14400 s 的拟合值。附件2半径采用 PCHIP 保单调插值，避免
产生非物理回弹。
"""
from __future__ import annotations

from pathlib import Path
import sys

import numpy as np
from openpyxl import load_workbook
from scipy.interpolate import PchipInterpolator

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src" / "q2_full"))
import solver_q2 as q2


ATTACHMENT2 = ROOT / "data" / "raw" / "附件2.xlsx"
THRESHOLD = 0.15
MAX_TIME = 400000.0
BOUNDARY_DATA_END = 14400.0
C_FLOOR = 1.0e-6


def read_radius_curve(path: Path = ATTACHMENT2):
    """读取附件2（s, cm），返回以 m 为单位的 PCHIP 半径函数及其导数。"""
    ws = load_workbook(path, data_only=True, read_only=True).active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    time = np.asarray([float(row[0]) for row in rows if row[0] is not None], dtype=float)
    radius_cm = np.asarray([float(row[1]) for row in rows if row[0] is not None], dtype=float)
    if time.size < 2 or np.any(np.diff(time) <= 0) or np.any(radius_cm <= 0):
        raise ValueError("附件2的时间或半径数据不满足插值条件")
    spline = PchipInterpolator(time, radius_cm * 0.01, extrapolate=False)
    derivative = spline.derivative()

    def radius(t):
        t = np.clip(np.asarray(t, dtype=float), time[0], time[-1])
        return spline(t)

    def radius_rate(t):
        t = np.clip(np.asarray(t, dtype=float), time[0], time[-1])
        rate = derivative(t)
        # 附件2末段半径恒定；端点微小数值误差不应造成虚假膨胀。
        return np.minimum(rate, 0.0)

    return time, radius_cm * 0.01, radius, radius_rate


def props_appendix4(c, t_c):
    """附录4经验式；扩散式中的温度严格采用开尔文。"""
    c = np.maximum(np.asarray(c, dtype=float), C_FLOOR)
    tk = np.asarray(t_c, dtype=float) + 273.15
    rho = 760.0 + 90.0 * c
    cp = 1850.0 + 2150.0 * c / (c + 1.0)
    k = 0.12 + 0.20 * c / (c + 1.0)
    d = 4.2e-4 * np.exp(-0.30 / c) * np.exp(-3850.0 / tk)
    return rho, cp, k, d


def ambient_temperature(t):
    """0--7200 s 取拟合曲线，此后固定为恒温设定值（摄氏度）。"""
    t = np.asarray(t, dtype=float)
    return np.where(t <= q2.T_SWITCH, q2.env_temperature(t), q2.T_PLATEAU_C)


def ambient_moisture(t):
    """附件1窗口内取拟合曲线，4 h 后固定为平台值。"""
    t = np.asarray(t, dtype=float)
    plateau = float(q2.env_moisture(BOUNDARY_DATA_END))
    return np.where(t <= BOUNDARY_DATA_END, q2.env_moisture(t), plateau)


def reference_nodes(refine=1):
    """沿用问题二的表面加密策略，并缩放到 xi∈[0,1]。"""
    nodes, _ = q2.build_nodes(refine)
    return nodes / q2.RADIUS


def assemble_moving_rate(nodes, diffusivity, storage, boundary_transfer, radius_m):
    """组装物质参考坐标中的扩散--收缩算子。

    xi 随收缩材料点运动，干基含水率以绝干物质为基准，故不另加欧拉坐标
    变换中出现的平流项；收缩通过 R(t) 对扩散项 W/R^2 和边界项 h/R 的
    共同缩放进入方程。
    """
    tri_full = q2.assemble_rate(
        nodes,
        np.asarray(diffusivity, dtype=float) / radius_m ** 2,
        storage,
        boundary_transfer / radius_m,
    )
    return tri_full


def interpolate_physical_profile(nodes_xi, profile, radius_m, distances_cm):
    """将固定 xi 解映射到指定物理半径；药材外部以 NaN 标记。"""
    physical_r = np.asarray(distances_cm, dtype=float) * 0.01
    valid = physical_r <= radius_m + 1e-12
    result = np.full(physical_r.size, np.nan)
    result[valid] = np.interp(physical_r[valid] / radius_m, nodes_xi, profile)
    return result


def run(refine=1, dt=10.0, save_interval=60.0, relaxation=0.7, tolerance=1.0e-9):
    nodes = reference_nodes(refine)
    n = nodes.size
    _, radius_data, radius, radius_rate = read_radius_curve()
    if not np.all(np.diff(radius_data) <= 1e-12):
        raise ValueError("附件2半径序列应为非增；请先核查数据")

    temperature = np.full(n, q2.T0_C)
    moisture = np.full(n, q2.C0)
    saved_t = [0.0]
    saved_T = [temperature.copy()]
    saved_C = [moisture.copy()]
    saved_R = [float(radius(0.0))]
    next_save = save_interval
    crossing = None
    max_iterations = 0

    steps = int(np.ceil(MAX_TIME / dt))
    for step in range(1, steps + 1):
        t_old = (step - 1) * dt
        t_new = step * dt
        r_new = float(radius(t_new))
        ta_old = q2.T0_C if step == 1 else float(ambient_temperature(t_old))
        ca_old = float(ambient_moisture(t_old))
        ta_new = float(ambient_temperature(t_new))
        ca_new = float(ambient_moisture(t_new))
        near_switch = (
            q2.T_SWITCH < t_new <= q2.T_SWITCH + 2.0 * dt
            or BOUNDARY_DATA_END < t_new <= BOUNDARY_DATA_END + 2.0 * dt
        )
        theta = 1.0 if step <= 2 or near_switch else 0.5
        old_max = float(np.max(moisture))

        t_guess = temperature.copy()
        c_guess = moisture.copy()
        for iteration in range(1, 81):
            rho, cp, k, _ = props_appendix4(c_guess, t_guess)
            rate_t = assemble_moving_rate(nodes, k, rho * cp, q2.H_T, r_new)
            t_field = q2.implicit_theta(temperature, rate_t[:3], rate_t[3], ta_new, ta_old, dt, theta)

            _, _, _, d = props_appendix4(c_guess, t_field)
            rate_c = assemble_moving_rate(nodes, d, np.ones(n), q2.H_M, r_new)
            c_field = q2.implicit_theta(moisture, rate_c[:3], rate_c[3], ca_new, ca_old, dt, theta)
            c_field = np.maximum(c_field, 0.0)
            error = max(
                np.max(np.abs(t_field - t_guess)) / (1.0 + np.max(np.abs(t_field))),
                np.max(np.abs(c_field - c_guess)) / (1.0 + np.max(np.abs(c_field))),
            )
            t_guess = relaxation * t_field + (1.0 - relaxation) * t_guess
            c_guess = relaxation * c_field + (1.0 - relaxation) * c_guess
            if error < tolerance:
                break
        else:
            raise RuntimeError(f"t={t_new}: Picard迭代未收敛")

        max_iterations = max(max_iterations, iteration)
        temperature, moisture = t_guess, c_guess
        new_max = float(np.max(moisture))
        if crossing is None and old_max >= THRESHOLD and new_max < THRESHOLD:
            fraction = (old_max - THRESHOLD) / (old_max - new_max)
            crossing = t_old + fraction * dt

        if t_new + 1e-9 >= next_save:
            saved_t.append(t_new)
            saved_T.append(temperature.copy())
            saved_C.append(moisture.copy())
            saved_R.append(r_new)
            next_save += save_interval
        if crossing is not None and t_new >= np.ceil(crossing / save_interval) * save_interval:
            break

    if crossing is None:
        raise RuntimeError(f"{MAX_TIME:g} s 内未达到干燥判据")
    return {
        "t": np.asarray(saved_t), "T": np.asarray(saved_T), "C": np.asarray(saved_C),
        "R": np.asarray(saved_R), "xi": nodes, "crossing": float(crossing),
        "max_iterations": max_iterations,
    }


def validate(result):
    c = result["C"]
    t = result["T"]
    r = result["R"]
    return {
        "finite": bool(np.isfinite(c).all() and np.isfinite(t).all()),
        "moisture_nonnegative": bool(c.min() >= -1e-10),
        "radius_monotone": bool(np.all(np.diff(r) <= 1e-12)),
        "temperature_within_boundary_range": bool(t.min() >= q2.T0_C - 1e-8 and t.max() <= ambient_temperature(BOUNDARY_DATA_END) + 1e-5),
        "final_moisture_radial_monotone": bool(np.all(np.diff(c[-1]) <= 1e-8)),
    }


if __name__ == "__main__":
    refine = int(sys.argv[1]) if len(sys.argv) > 1 else 1
    dt = float(sys.argv[2]) if len(sys.argv) > 2 else 10.0
    output = Path(sys.argv[3]) if len(sys.argv) > 3 else ROOT / "outputs/q4/q4_appendix4_shrink.npz"
    output.parent.mkdir(parents=True, exist_ok=True)
    result = run(refine=refine, dt=dt)
    checks = validate(result)
    np.savez(output, **result)
    print("crossing_s", result["crossing"])
    print("crossing_h", result["crossing"] / 3600.0)
    print("max_iterations", result["max_iterations"])
    for key, value in checks.items():
        print(key, value)
