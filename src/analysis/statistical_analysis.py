# -*- coding: utf-8 -*-
"""
药材烘干数据分析与拟合报告
====================================================
功能：
1. 附件1、附件2的描述性统计分析
2. 多种拟合模型对比（指数、多项式、对数等）
3. 残差分析与可视化
4. 相关性热力图
5. 完整的统计报告输出
"""

import matplotlib
matplotlib.use('Agg')  # 使用非交互式后端

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import seaborn as sns
from scipy import stats
from scipy.optimize import curve_fit
import openpyxl
import warnings
warnings.filterwarnings('ignore')

# 设置中文字体
plt.rcParams['font.sans-serif'] = ['SimHei', 'Microsoft YaHei']
plt.rcParams['axes.unicode_minus'] = False

# 设置绘图风格
sns.set_style("whitegrid")
sns.set_palette("husl")


class DataAnalyzer:
    """数据分析器"""

    def __init__(self, file_path):
        self.file_path = file_path
        self.data = None
        self.load_data()

    def load_data(self):
        """加载Excel数据"""
        wb = openpyxl.load_workbook(self.file_path, data_only=True)
        ws = wb.active
        rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]

        if len(rows[0]) == 3:  # 附件1格式
            self.data = pd.DataFrame(rows, columns=['时间(s)', '温度(°C)', '水分含量(kg/kg)'])
        else:  # 附件2格式
            cols = [c.value for c in list(ws.iter_rows(min_row=1, max_row=1))[0] if c.value]
            self.data = pd.DataFrame(rows, columns=cols)

        print(f"✓ 已加载数据: {self.file_path}")
        print(f"  数据形状: {self.data.shape}")

    def descriptive_stats(self):
        """描述性统计分析"""
        print("\n" + "="*70)
        print("描述性统计分析")
        print("="*70)

        stats_df = pd.DataFrame()

        for col in self.data.select_dtypes(include=[np.number]).columns:
            data = self.data[col].dropna()

            stats_dict = {
                '变量': col,
                '样本数': len(data),
                '均值': data.mean(),
                '标准差': data.std(),
                '方差': data.var(),
                '最小值': data.min(),
                '25%分位': data.quantile(0.25),
                '中位数': data.median(),
                '75%分位': data.quantile(0.75),
                '最大值': data.max(),
                '极差': data.max() - data.min(),
                '偏度': stats.skew(data),
                '峰度': stats.kurtosis(data),
                '变异系数': data.std() / data.mean() if data.mean() != 0 else 0
            }

            stats_df = pd.concat([stats_df, pd.DataFrame([stats_dict])], ignore_index=True)

        print(stats_df.to_string(index=False))

        # 偏度和峰度解释
        print("\n" + "-"*70)
        print("统计量解释：")
        print("-"*70)
        for _, row in stats_df.iterrows():
            print(f"\n【{row['变量']}】")

            # 偏度解释
            skewness = row['偏度']
            if abs(skewness) < 0.5:
                skew_desc = "近似对称分布"
            elif skewness > 0:
                skew_desc = f"右偏分布（正偏，偏度={skewness:.3f}）"
            else:
                skew_desc = f"左偏分布（负偏，偏度={skewness:.3f}）"
            print(f"  偏度: {skew_desc}")

            # 峰度解释
            kurtosis = row['峰度']
            if abs(kurtosis) < 0.5:
                kurt_desc = "正态峰度（mesokurtic）"
            elif kurtosis > 0:
                kurt_desc = f"尖峰分布（leptokurtic，峰度={kurtosis:.3f}）- 尾部较厚"
            else:
                kurt_desc = f"平峰分布（platykurtic，峰度={kurtosis:.3f}）- 尾部较薄"
            print(f"  峰度: {kurt_desc}")

            # 变异系数
            cv = row['变异系数']
            if cv < 0.1:
                cv_desc = "低变异性（CV<10%）"
            elif cv < 0.3:
                cv_desc = "中等变异性（10%<CV<30%）"
            else:
                cv_desc = f"高变异性（CV={cv*100:.1f}%）"
            print(f"  变异系数: {cv_desc}")

        return stats_df

    def normality_test(self):
        """正态性检验"""
        print("\n" + "="*70)
        print("正态性检验 (Shapiro-Wilk Test)")
        print("="*70)

        for col in self.data.select_dtypes(include=[np.number]).columns:
            data = self.data[col].dropna()

            if len(data) >= 3:
                stat, p_value = stats.shapiro(data)
                result = "服从正态分布" if p_value > 0.05 else "不服从正态分布"
                print(f"{col:20s}: W={stat:.4f}, p={p_value:.6f} → {result}")

    def plot_distributions(self, save_path=None):
        """绘制数据分布图"""
        numeric_cols = self.data.select_dtypes(include=[np.number]).columns
        n_cols = len(numeric_cols)

        fig, axes = plt.subplots(n_cols, 2, figsize=(14, 4*n_cols))
        if n_cols == 1:
            axes = axes.reshape(1, -1)

        for idx, col in enumerate(numeric_cols):
            data = self.data[col].dropna()

            # 直方图 + KDE
            axes[idx, 0].hist(data, bins=30, density=True, alpha=0.6, color='skyblue', edgecolor='black')
            data.plot.kde(ax=axes[idx, 0], color='red', linewidth=2)
            axes[idx, 0].set_title(f'{col} - 分布图', fontsize=12, fontweight='bold')
            axes[idx, 0].set_xlabel(col)
            axes[idx, 0].set_ylabel('概率密度')
            axes[idx, 0].grid(alpha=0.3)

            # Q-Q图（正态性检验）
            stats.probplot(data, dist="norm", plot=axes[idx, 1])
            axes[idx, 1].set_title(f'{col} - Q-Q图（正态性）', fontsize=12, fontweight='bold')
            axes[idx, 1].grid(alpha=0.3)

        plt.tight_layout()
        if save_path:
            plt.savefig(save_path, dpi=300, bbox_inches='tight')
            print(f"✓ 分布图已保存: {save_path}")
        plt.close()


class FittingAnalyzer:
    """拟合分析器"""

    def __init__(self, x, y, var_name="变量"):
        self.x = np.array(x)
        self.y = np.array(y)
        self.var_name = var_name
        self.results = {}

    @staticmethod
    def exponential_model(x, a, b, c):
        """指数模型: y = a + b*exp(c*x)"""
        return a + b * np.exp(c * x)

    @staticmethod
    def double_exponential_model(x, a, b1, c1, b2, c2):
        """双指数模型: y = a + b1*exp(c1*x) + b2*exp(c2*x)"""
        return a + b1 * np.exp(c1 * x) + b2 * np.exp(c2 * x)

    @staticmethod
    def polynomial_model_2(x, a, b, c):
        """二次多项式: y = a + b*x + c*x^2"""
        return a + b*x + c*x**2

    @staticmethod
    def polynomial_model_3(x, a, b, c, d):
        """三次多项式: y = a + b*x + c*x^2 + d*x^3"""
        return a + b*x + c*x**2 + d*x**3

    @staticmethod
    def logarithmic_model(x, a, b):
        """对数模型: y = a + b*ln(x+1)"""
        return a + b * np.log(x + 1)

    @staticmethod
    def power_model(x, a, b):
        """幂函数模型: y = a * x^b"""
        return a * x**b

    def fit_model(self, model_func, p0=None, model_name="Model"):
        """拟合模型并计算统计量"""
        try:
            # 拟合
            popt, pcov = curve_fit(model_func, self.x, self.y, p0=p0, maxfev=10000)
            y_pred = model_func(self.x, *popt)

            # 残差
            residuals = self.y - y_pred
            ss_res = np.sum(residuals**2)
            ss_tot = np.sum((self.y - np.mean(self.y))**2)

            # 统计量
            r2 = 1 - (ss_res / ss_tot)
            rmse = np.sqrt(np.mean(residuals**2))
            mae = np.mean(np.abs(residuals))
            mape = np.mean(np.abs(residuals / self.y)) * 100 if np.all(self.y != 0) else np.inf

            # 调整R²
            n = len(self.y)
            k = len(popt)
            adj_r2 = 1 - (1 - r2) * (n - 1) / (n - k - 1)

            # AIC & BIC
            aic = n * np.log(ss_res / n) + 2 * k
            bic = n * np.log(ss_res / n) + k * np.log(n)

            self.results[model_name] = {
                'params': popt,
                'covariance': pcov,
                'y_pred': y_pred,
                'residuals': residuals,
                'R²': r2,
                'Adj_R²': adj_r2,
                'RMSE': rmse,
                'MAE': mae,
                'MAPE': mape,
                'AIC': aic,
                'BIC': bic,
                'model_func': model_func
            }

            return True
        except Exception as e:
            print(f"✗ {model_name} 拟合失败: {e}")
            return False

    def fit_all_models(self):
        """拟合所有模型"""
        print(f"\n正在拟合 {self.var_name} ...")

        # 指数模型
        self.fit_model(self.exponential_model, p0=[self.y[0], self.y[-1]-self.y[0], 0.001],
                      model_name="指数模型")

        # 双指数模型
        self.fit_model(self.double_exponential_model,
                      p0=[self.y[0], (self.y[-1]-self.y[0])/2, 0.001, (self.y[-1]-self.y[0])/2, -0.001],
                      model_name="双指数模型")

        # 多项式模型
        self.fit_model(self.polynomial_model_2, model_name="二次多项式")
        self.fit_model(self.polynomial_model_3, model_name="三次多项式")

        # 对数模型
        if np.all(self.x >= 0):
            self.fit_model(self.logarithmic_model, p0=[self.y[0], 1], model_name="对数模型")

        # 幂函数模型
        if np.all(self.x > 0) and np.all(self.y > 0):
            self.fit_model(self.power_model, p0=[self.y[0], 0.5], model_name="幂函数模型")

    def print_comparison(self):
        """打印模型对比"""
        print("\n" + "="*100)
        print(f"{self.var_name} - 拟合模型对比")
        print("="*100)

        comparison_df = pd.DataFrame({
            '模型': list(self.results.keys()),
            'R²': [self.results[m]['R²'] for m in self.results],
            'Adj_R²': [self.results[m]['Adj_R²'] for m in self.results],
            'RMSE': [self.results[m]['RMSE'] for m in self.results],
            'MAE': [self.results[m]['MAE'] for m in self.results],
            'MAPE(%)': [self.results[m]['MAPE'] for m in self.results],
            'AIC': [self.results[m]['AIC'] for m in self.results],
            'BIC': [self.results[m]['BIC'] for m in self.results]
        })

        # 排序
        comparison_df = comparison_df.sort_values('R²', ascending=False)

        print(comparison_df.to_string(index=False))

        # 最佳模型
        best_model = comparison_df.iloc[0]['模型']
        print(f"\n✓ 最佳模型（基于R²）: {best_model}")
        print(f"  R² = {self.results[best_model]['R²']:.6f}")
        print(f"  RMSE = {self.results[best_model]['RMSE']:.6f}")

        # 参数
        print(f"\n模型参数:")
        params = self.results[best_model]['params']
        for i, p in enumerate(params):
            print(f"  参数{i+1}: {p:.10f}")

        return comparison_df

    def residual_analysis(self, model_name=None):
        """残差分析"""
        if model_name is None:
            # 选择R²最高的模型
            model_name = max(self.results.keys(), key=lambda k: self.results[k]['R²'])

        if model_name not in self.results:
            print(f"✗ 模型 {model_name} 不存在")
            return

        residuals = self.results[model_name]['residuals']

        print("\n" + "="*70)
        print(f"{self.var_name} - 残差分析 ({model_name})")
        print("="*70)

        # 残差统计
        print(f"残差均值: {np.mean(residuals):.6e} (应接近0)")
        print(f"残差标准差: {np.std(residuals):.6f}")
        print(f"残差最大值: {np.max(residuals):.6f}")
        print(f"残差最小值: {np.min(residuals):.6f}")

        # 正态性检验
        if len(residuals) >= 3:
            stat, p_value = stats.shapiro(residuals)
            result = "服从正态分布" if p_value > 0.05 else "不服从正态分布"
            print(f"Shapiro-Wilk检验: W={stat:.4f}, p={p_value:.6f} → 残差{result}")

        # Durbin-Watson检验（自相关）
        dw = np.sum(np.diff(residuals)**2) / np.sum(residuals**2)
        print(f"Durbin-Watson统计量: {dw:.4f} (接近2表示无自相关)")

    def plot_fitting_results(self, save_path=None):
        """绘制拟合结果"""
        n_models = len(self.results)

        fig, axes = plt.subplots(2, n_models, figsize=(5*n_models, 10))
        if n_models == 1:
            axes = axes.reshape(-1, 1)

        for idx, (model_name, result) in enumerate(self.results.items()):
            y_pred = result['y_pred']
            residuals = result['residuals']

            # 拟合曲线
            axes[0, idx].scatter(self.x, self.y, alpha=0.6, s=30, label='原始数据')
            axes[0, idx].plot(self.x, y_pred, 'r-', linewidth=2, label='拟合曲线')
            axes[0, idx].set_title(f'{model_name}\nR²={result["R²"]:.6f}', fontsize=11, fontweight='bold')
            axes[0, idx].set_xlabel('时间 (s)')
            axes[0, idx].set_ylabel(self.var_name)
            axes[0, idx].legend()
            axes[0, idx].grid(alpha=0.3)

            # 残差图
            axes[1, idx].scatter(self.x, residuals, alpha=0.6, s=30, color='purple')
            axes[1, idx].axhline(y=0, color='r', linestyle='--', linewidth=2)
            axes[1, idx].set_title(f'残差图\nRMSE={result["RMSE"]:.6f}', fontsize=11, fontweight='bold')
            axes[1, idx].set_xlabel('时间 (s)')
            axes[1, idx].set_ylabel('残差')
            axes[1, idx].grid(alpha=0.3)

        plt.tight_layout()
        if save_path:
            plt.savefig(save_path, dpi=300, bbox_inches='tight')
            print(f"✓ 拟合图已保存: {save_path}")
        plt.close()


def correlation_analysis(data, save_path=None):
    """相关性分析与热力图"""
    print("\n" + "="*70)
    print("相关性分析 (Pearson相关系数)")
    print("="*70)

    numeric_data = data.select_dtypes(include=[np.number])
    corr_matrix = numeric_data.corr()

    print(corr_matrix)

    # 热力图
    plt.figure(figsize=(10, 8))
    sns.heatmap(corr_matrix, annot=True, fmt='.4f', cmap='coolwarm',
                center=0, square=True, linewidths=1, cbar_kws={"shrink": 0.8})
    plt.title('相关性热力图', fontsize=16, fontweight='bold', pad=20)
    plt.tight_layout()

    if save_path:
        plt.savefig(save_path, dpi=300, bbox_inches='tight')
        print(f"✓ 热力图已保存: {save_path}")
    plt.close()

    return corr_matrix


def main():
    """主函数"""
    print("="*70)
    print("药材烘干数据分析报告")
    print("="*70)

    # 创建输出目录
    import os
    output_dir = "outputs/data_analysis"
    os.makedirs(output_dir, exist_ok=True)

    # ========== 附件1分析 ==========
    print("\n" + "#"*70)
    print("# 附件1：烘房环境温度和水分含量")
    print("#"*70)

    analyzer1 = DataAnalyzer("../../data/raw/附件1.xlsx")

    # 描述性统计
    stats1 = analyzer1.descriptive_stats()
    stats1.to_excel(f"{output_dir}/附件1_描述性统计.xlsx", index=False)

    # 正态性检验
    analyzer1.normality_test()

    # 分布图
    analyzer1.plot_distributions(save_path=f"{output_dir}/附件1_分布图.png")

    # 相关性
    corr1 = correlation_analysis(analyzer1.data, save_path=f"{output_dir}/附件1_相关性热力图.png")

    # ========== 拟合分析 ==========
    t = analyzer1.data['时间(s)'].values
    temp = analyzer1.data['温度(°C)'].values
    moisture = analyzer1.data['水分含量(kg/kg)'].values

    # 温度拟合
    print("\n" + "#"*70)
    print("# 温度拟合分析")
    print("#"*70)
    fitter_temp = FittingAnalyzer(t, temp, var_name="温度(°C)")
    fitter_temp.fit_all_models()
    comp_temp = fitter_temp.print_comparison()
    comp_temp.to_excel(f"{output_dir}/温度拟合对比.xlsx", index=False)
    fitter_temp.residual_analysis()
    fitter_temp.plot_fitting_results(save_path=f"{output_dir}/温度拟合结果.png")

    # 水分拟合
    print("\n" + "#"*70)
    print("# 水分含量拟合分析")
    print("#"*70)
    fitter_moisture = FittingAnalyzer(t, moisture, var_name="水分含量(kg/kg)")
    fitter_moisture.fit_all_models()
    comp_moisture = fitter_moisture.print_comparison()
    comp_moisture.to_excel(f"{output_dir}/水分拟合对比.xlsx", index=False)
    fitter_moisture.residual_analysis()
    fitter_moisture.plot_fitting_results(save_path=f"{output_dir}/水分拟合结果.png")

    # ========== 附件2分析 ==========
    print("\n" + "#"*70)
    print("# 附件2：物性参数")
    print("#"*70)

    analyzer2 = DataAnalyzer("../../data/raw/附件2.xlsx")
    stats2 = analyzer2.descriptive_stats()
    stats2.to_excel(f"{output_dir}/附件2_描述性统计.xlsx", index=False)
    analyzer2.normality_test()
    analyzer2.plot_distributions(save_path=f"{output_dir}/附件2_分布图.png")

    if len(analyzer2.data.select_dtypes(include=[np.number]).columns) > 1:
        corr2 = correlation_analysis(analyzer2.data, save_path=f"{output_dir}/附件2_相关性热力图.png")

    print("\n" + "="*70)
    print("✓ 所有分析完成！结果已保存到:", output_dir)
    print("="*70)


if __name__ == "__main__":
    main()
